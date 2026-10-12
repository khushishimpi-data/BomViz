import datetime as dt
import re
import shutil
import time
import zipfile
from collections.abc import Callable
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from app.services.mapping_suggestions import count_standard_field_matches, suggest_mapping
from app.services.parsing import (
    ParseError,
    ParseErrorCode,
    ParseLimits,
    ParseOptions,
    ParseOutcome,
    parse_file,
)

FIXTURES = Path(__file__).resolve().parents[2] / "fixtures" / "boms"
SHEET1 = "xl/worksheets/sheet1.xml"


def save(workbook: Workbook, tmp_path: Path, name: str = "bom.xlsx") -> Path:
    path = tmp_path / name
    workbook.save(path)
    return path


def rewrite_member(path: Path, member: str, transform: Callable[[str], str]) -> None:
    with zipfile.ZipFile(path) as source:
        items = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in items:
            target.writestr(name, transform(data.decode()).encode() if name == member else data)


def bom_workbook() -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "BOM"
    sheet["A1"] = "Assembly BOM"
    sheet.append([])
    sheet.append(["Part Number", "Qty", "Unit Cost", "Active", "Released", "Updated"])
    sheet.append(["P-001", 4, 0.1, True, dt.date(2026, 10, 11), dt.datetime(2026, 10, 11, 8, 30)])
    sheet.append(["P-002", 2.5, 1234.5, False, None, None])
    return workbook


def expect_error(path: Path, code: ParseErrorCode, limits: ParseLimits | None = None) -> None:
    with pytest.raises(ParseError) as exc:
        parse_file(path, "XLSX", limits=limits)
    assert exc.value.code is code


def test_parses_typed_cells_below_a_title_row(tmp_path: Path) -> None:
    outcome = parse_file(save(bom_workbook(), tmp_path), "XLSX")

    assert outcome.effective_options == ParseOptions(sheet_name="BOM", header_row=3)
    assert [c.name for c in outcome.columns][:2] == ["Part Number", "Qty"]
    assert outcome.row_count == 2
    assert [(r.row_number, r.cells) for r in outcome.sample_rows] == [
        (4, ["P-001", "4", "0.1", "TRUE", "2026-10-11", "2026-10-11T08:30:00"]),
        (5, ["P-002", "2.5", "1234.5", "FALSE", "", ""]),
    ]


def test_dates_in_a_1904_workbook_are_read_correctly(tmp_path: Path) -> None:
    workbook = bom_workbook()
    workbook.epoch = CALENDAR_MAC_1904

    outcome = parse_file(save(workbook, tmp_path), "XLSX")

    assert outcome.sample_rows[0].cells[4] == "2026-10-11"


def test_formulas_yield_their_cached_value_or_empty(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Part", "Qty", "Double"])
    sheet.append(["A", 2, "=B2*2"])
    sheet.append(["B", 3, "=B3*2"])
    path = save(workbook, tmp_path)
    rewrite_member(path, SHEET1, lambda xml: xml.replace("<f>B2*2</f><v />", "<f>B2*2</f><v>4</v>"))

    outcome = parse_file(path, "XLSX")

    assert [r.cells for r in outcome.sample_rows] == [["A", "2", "4"], ["B", "3", ""]]


def test_rows_beyond_a_wrong_declared_dimension_are_read(tmp_path: Path) -> None:
    path = save(bom_workbook(), tmp_path)
    rewrite_member(
        path, SHEET1, lambda xml: re.sub(r'<dimension ref="[^"]*"', '<dimension ref="A1"', xml)
    )

    outcome = parse_file(path, "XLSX")

    assert outcome.row_count == 2
    assert outcome.column_count == 6


def multi_sheet_workbook() -> Workbook:
    workbook = Workbook()
    notes = workbook.active
    assert notes is not None
    notes.title = "Notes"
    notes.append(["Hidden A", "Hidden B"])
    notes.append(["x", "y"])
    notes.sheet_state = "hidden"
    workbook.create_sheet("Blank")
    parts = workbook.create_sheet("Parts")
    parts.append(["Part", "Qty"])
    parts.append(["A-1", 3])
    return workbook


def test_default_sheet_is_the_first_visible_sheet_with_data(tmp_path: Path) -> None:
    outcome = parse_file(save(multi_sheet_workbook(), tmp_path), "XLSX")

    assert outcome.effective_options.sheet_name == "Parts"
    assert [(s.name, s.visible) for s in outcome.sheets] == [
        ("Notes", False),
        ("Blank", True),
        ("Parts", True),
    ]
    assert outcome.sample_rows[0].cells == ["A-1", "3"]


def test_hidden_sheet_can_be_chosen_explicitly(tmp_path: Path) -> None:
    path = save(multi_sheet_workbook(), tmp_path)

    outcome = parse_file(path, "XLSX", ParseOptions(sheet_name="Notes"))

    assert [c.name for c in outcome.columns] == ["Hidden A", "Hidden B"]


def test_unknown_sheet_fails(tmp_path: Path) -> None:
    path = save(multi_sheet_workbook(), tmp_path)

    with pytest.raises(ParseError) as exc:
        parse_file(path, "XLSX", ParseOptions(sheet_name="Missing"))

    assert exc.value.code is ParseErrorCode.SHEET_NOT_FOUND


def test_explicitly_chosen_empty_sheet_is_empty(tmp_path: Path) -> None:
    path = save(multi_sheet_workbook(), tmp_path)

    with pytest.raises(ParseError) as exc:
        parse_file(path, "XLSX", ParseOptions(sheet_name="Blank"))

    assert exc.value.code is ParseErrorCode.EMPTY_FILE


def test_sheet_list_is_capped(tmp_path: Path) -> None:
    path = save(multi_sheet_workbook(), tmp_path)

    outcome = parse_file(path, "XLSX", limits=ParseLimits(max_sheets_listed=2))

    assert [s.name for s in outcome.sheets] == ["Notes", "Blank"]
    assert outcome.effective_options.sheet_name == "Parts"


def test_workbook_without_data_is_empty(tmp_path: Path) -> None:
    expect_error(save(Workbook(), tmp_path), ParseErrorCode.EMPTY_FILE)


def test_row_limit_applies_to_workbooks(tmp_path: Path) -> None:
    expect_error(
        save(bom_workbook(), tmp_path), ParseErrorCode.TOO_MANY_ROWS, ParseLimits(max_rows=1)
    )


def test_member_with_extreme_compression_ratio_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "bomb.xlsx"
    chunk = bytes(64 * 1024)
    with (
        zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive,
        archive.open(SHEET1, "w") as member,
    ):
        for _ in range(64):  # 4 MiB of zeros compresses to a few KiB
            member.write(chunk)

    expect_error(path, ParseErrorCode.FILE_TOO_COMPLEX)


def test_total_uncompressed_size_over_the_limit_is_rejected(tmp_path: Path) -> None:
    path = save(bom_workbook(), tmp_path)

    expect_error(path, ParseErrorCode.FILE_TOO_COMPLEX, ParseLimits(max_uncompressed_bytes=1024))


def test_too_many_archive_members_are_rejected(tmp_path: Path) -> None:
    path = save(bom_workbook(), tmp_path)

    expect_error(path, ParseErrorCode.FILE_TOO_COMPLEX, ParseLimits(max_zip_members=3))


def test_bytes_that_are_not_a_zip_are_corrupt(tmp_path: Path) -> None:
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"PK\x03\x04" + b"\x00garbage" * 100)

    expect_error(path, ParseErrorCode.FILE_CORRUPT)


def test_truncated_workbook_is_corrupt(tmp_path: Path) -> None:
    path = save(bom_workbook(), tmp_path)
    path.write_bytes(path.read_bytes()[:2000])

    expect_error(path, ParseErrorCode.FILE_CORRUPT)


def test_zip_that_is_not_a_workbook_is_corrupt(tmp_path: Path) -> None:
    path = tmp_path / "other.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("readme.txt", "hello")

    expect_error(path, ParseErrorCode.FILE_CORRUPT)


def test_ole_file_named_xlsx_is_reported_as_encrypted(tmp_path: Path) -> None:
    path = tmp_path / "protected.xlsx"
    shutil.copyfile(FIXTURES / "basic.xls", path)

    expect_error(path, ParseErrorCode.ENCRYPTED_WORKBOOK)


# Crafted sheets: a valid workbook whose sheet XML is replaced by hand-written rows, the way
# a hostile producer would write it. Each must fail quickly, whatever the row numbers say.
WALL_CLOCK_LIMIT_SECONDS = 3.0


def inline_cell(ref: str, text: str) -> str:
    return f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>'


def crafted_workbook(tmp_path: Path, sheet_data: str) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "placeholder"
    path = save(workbook, tmp_path, "crafted.xlsx")
    rewrite_member(
        path,
        SHEET1,
        lambda xml: re.sub(
            r"<sheetData>.*</sheetData>", f"<sheetData>{sheet_data}</sheetData>", xml, flags=re.S
        ),
    )
    return path


def expect_fast_error(path: Path, code: ParseErrorCode, limits: ParseLimits | None = None) -> None:
    started = time.perf_counter()
    expect_error(path, code, limits)
    assert time.perf_counter() - started < WALL_CLOCK_LIMIT_SECONDS


@pytest.mark.parametrize("row", [5_000_000, 900_000_000])
def test_single_row_far_down_the_sheet_fails_fast(tmp_path: Path, row: int) -> None:
    cells = inline_cell(f"A{row}", "Part") + inline_cell(f"B{row}", "Qty")
    path = crafted_workbook(tmp_path, f'<row r="{row}">{cells}</row>')

    expect_fast_error(path, ParseErrorCode.FILE_TOO_COMPLEX)


def test_many_rows_padded_to_the_last_column_fail_fast(tmp_path: Path) -> None:
    rows = "".join(f'<row r="{i}"><c r="ZZZ{i}"/></row>' for i in range(1, 20_001))
    path = crafted_workbook(tmp_path, rows)

    expect_fast_error(path, ParseErrorCode.FILE_TOO_COMPLEX)


def test_huge_gap_below_the_header_fails_fast(tmp_path: Path) -> None:
    rows = (
        f'<row r="1">{inline_cell("A1", "Part")}{inline_cell("B1", "Qty")}</row>'
        f'<row r="2">{inline_cell("A2", "P-1")}{inline_cell("B2", "4")}</row>'
        f'<row r="1000000">{inline_cell("A1000000", "P-2")}</row>'
    )
    path = crafted_workbook(tmp_path, rows)

    expect_fast_error(path, ParseErrorCode.FILE_TOO_COMPLEX)


def test_row_beyond_the_excel_maximum_fails_even_when_empty_rows_are_allowed(
    tmp_path: Path,
) -> None:
    row = 1_048_577
    path = crafted_workbook(tmp_path, f'<row r="{row}">{inline_cell(f"A{row}", "x")}</row>')
    lenient = ParseLimits(max_leading_empty_rows=10_000_000, max_skipped_empty_rows=10_000_000)

    expect_fast_error(path, ParseErrorCode.FILE_TOO_COMPLEX, lenient)


def test_row_at_the_excel_maximum_is_read(tmp_path: Path) -> None:
    last = 1_048_576
    rows = (
        f'<row r="1">{inline_cell("A1", "Part")}{inline_cell("B1", "Qty")}</row>'
        f'<row r="{last}">{inline_cell(f"A{last}", "P-1")}{inline_cell(f"B{last}", "4")}</row>'
    )
    path = crafted_workbook(tmp_path, rows)
    lenient = ParseLimits(max_skipped_empty_rows=10_000_000)

    outcome = parse_file(path, "XLSX", limits=lenient)

    assert [(r.row_number, r.cells) for r in outcome.sample_rows] == [(last, ["P-1", "4"])]


def test_value_in_the_column_after_the_limit_fails(tmp_path: Path) -> None:
    rows = (
        f'<row r="1">{inline_cell("A1", "Part")}{inline_cell("B1", "Qty")}</row>'
        f'<row r="2">{inline_cell("A2", "P-1")}{inline_cell("D2", "extra")}</row>'
    )
    path = crafted_workbook(tmp_path, rows)

    expect_error(path, ParseErrorCode.TOO_MANY_COLUMNS, ParseLimits(max_columns=3))


def test_empty_cells_far_to_the_right_are_ignored(tmp_path: Path) -> None:
    rows = (
        f'<row r="1">{inline_cell("A1", "Part")}{inline_cell("B1", "Qty")}<c r="ZZZ1"/></row>'
        f'<row r="2">{inline_cell("A2", "P-1")}{inline_cell("B2", "4")}<c r="ZZZ2"/></row>'
    )
    path = crafted_workbook(tmp_path, rows)

    outcome = parse_file(path, "XLSX")

    assert [c.name for c in outcome.columns] == ["Part", "Qty"]
    assert outcome.sample_rows[0].cells == ["P-1", "4"]


def test_members_that_together_compress_too_well_are_rejected(tmp_path: Path) -> None:
    path = tmp_path / "spread-bomb.xlsx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        # Each member stays below the 1 MiB per-member floor; together they exceed it.
        for index in range(3):
            archive.writestr(f"xl/worksheets/sheet{index + 1}.xml", bytes(900 * 1024))

    expect_error(path, ParseErrorCode.FILE_TOO_COMPLEX)


# --- manufacturing export: information block and never-calculated formulas ------------------

FAN_HEADER = [
    "Level",
    "Item Code",
    "Component / Material",
    "Specification / Notes",
    "UOM",
    "Qty per Finished Unit",
    "Scrap %",
    "Effective Qty per Unit",
    "Batch Qty Required",
    "Unit Cost (₹)",
    "Extended Cost per Unit (₹)",
    "Batch Cost (₹)",
]
FAN_DATA_ROWS = range(11, 25)


def fan_workbook() -> Workbook:
    """The shape of a manufacturing BOM export whose formulas were never calculated."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Manufacturing BOM"
    sheet["A1"] = "MANUFACTURING BILL OF MATERIALS"
    info: list[tuple[str, str | float]] = [
        ("Product", "Electric Table Fan — 400 mm"),
        ("Finished Good Code", "FG-FAN-400"),
        ("BOM Revision", "A"),
        ("BOM Type", "Manufacturing BOM"),
        ("Production Batch Size (units)", 100),
        ("Currency", "INR (₹)"),
    ]
    for row, (key, value) in enumerate(info, start=3):
        sheet.cell(row=row, column=1, value=key)
        sheet.cell(row=row, column=2, value=value)
    for column, name in enumerate(FAN_HEADER, start=1):
        sheet.cell(row=10, column=column, value=name)
    for row in FAN_DATA_ROWS:
        values: list[str | float] = [
            1,
            f"RM-{row - 10:03d}",
            f"Part {row}",
            "Steel",
            "Nos",
            1,
            0.02,
        ]
        formulas: list[str | float] = [
            f"=F{row}*(1+G{row})",
            f"=H{row}*$B$7",
            85,
            f"=H{row}*J{row}",
        ]
        for column, value in enumerate([*values, *formulas, f"=I{row}*J{row}"], start=1):
            sheet.cell(row=row, column=column, value=value)
    sheet["C25"] = "TOTAL MATERIAL COST"
    return workbook


def parse_fan(path: Path) -> ParseOutcome:
    return parse_file(path, "XLSX", header_scorer=count_standard_field_matches)


def test_fan_export_detects_the_table_header_below_the_information_block(tmp_path: Path) -> None:
    outcome = parse_fan(save(fan_workbook(), tmp_path))

    assert outcome.effective_options == ParseOptions(sheet_name="Manufacturing BOM", header_row=10)
    assert [c.name for c in outcome.columns] == FAN_HEADER
    assert outcome.row_count == len(FAN_DATA_ROWS) + 1  # the TOTAL row is data in the raw stage


def test_fan_export_suggests_every_required_field(tmp_path: Path) -> None:
    outcome = parse_fan(save(fan_workbook(), tmp_path))

    mapping = suggest_mapping(outcome.columns)
    targets = {c.source_column_name: c.target_field for c in mapping.columns}
    assert targets["Level"] == "bom_level"
    assert targets["Item Code"] == "part_number"
    assert targets["Component / Material"] == "part_name"
    assert targets["UOM"] == "unit_of_measure"
    assert targets["Qty per Finished Unit"] == "quantity"
    assert targets["Unit Cost (₹)"] == "unit_cost"
    assert mapping.missing_required == []


def test_fan_export_counts_formula_cells_without_saved_values(tmp_path: Path) -> None:
    outcome = parse_fan(save(fan_workbook(), tmp_path))

    assert outcome.formula_cells_without_value == 4 * len(FAN_DATA_ROWS)


def test_formula_cells_with_cached_values_are_not_counted(tmp_path: Path) -> None:
    path = save(fan_workbook(), tmp_path)
    rewrite_member(path, SHEET1, lambda xml: xml.replace("<v />", "<v>1</v>"))

    outcome = parse_fan(path)

    assert outcome.formula_cells_without_value == 0
    assert outcome.sample_rows[0].cells[7] == "1"


def test_formula_cells_without_a_value_element_are_counted(tmp_path: Path) -> None:
    path = save(fan_workbook(), tmp_path)
    rewrite_member(path, SHEET1, lambda xml: xml.replace("<v />", ""))

    assert parse_fan(path).formula_cells_without_value == 4 * len(FAN_DATA_ROWS)


def test_formula_cells_above_the_header_row_are_not_counted(tmp_path: Path) -> None:
    workbook = bom_workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["B1"] = "=1+1"  # in the title row, above the header on row 3
    sheet["C4"] = "=B4*2"

    outcome = parse_file(save(workbook, tmp_path), "XLSX")

    assert outcome.effective_options.header_row == 3
    assert outcome.formula_cells_without_value == 1


def test_formula_count_follows_the_explicit_header_row(tmp_path: Path) -> None:
    path = save(fan_workbook(), tmp_path)

    outcome = parse_file(path, "XLSX", ParseOptions(header_row=20))

    assert outcome.formula_cells_without_value == 4 * len(range(20, 25))


def test_formula_count_reads_only_the_chosen_sheet(tmp_path: Path) -> None:
    workbook = fan_workbook()
    other = workbook.create_sheet("Plain")
    other.append(["Part Number", "Qty"])
    other.append(["P-1", 2])

    outcome = parse_file(save(workbook, tmp_path), "XLSX", ParseOptions(sheet_name="Plain"))

    assert outcome.formula_cells_without_value == 0


def test_workbook_without_formulas_counts_zero(tmp_path: Path) -> None:
    assert parse_file(save(bom_workbook(), tmp_path), "XLSX").formula_cells_without_value == 0
