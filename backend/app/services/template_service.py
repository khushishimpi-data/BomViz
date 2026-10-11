"""Downloadable BOM templates whose header row is the standard field labels.

Columns follow `STANDARD_FIELDS` order (required fields first), so a filled-in template parses
back and maps every column to its field with full confidence.
"""

import csv
import io
from types import MappingProxyType

from openpyxl import Workbook
from openpyxl.styles import Font

from app.domain.bom_fields import STANDARD_FIELDS

BOM_SHEET_TITLE = "BOM"
INSTRUCTIONS_SHEET_TITLE = "Instructions"

_FIELD_DESCRIPTIONS: MappingProxyType[str, str] = MappingProxyType(
    {
        "part_number": "Unique identifier of the part within this BOM.",
        "part_name": "Short name of the part.",
        "quantity": "Quantity used per parent assembly.",
        "parent_part_number": "Part Number of the assembly this part belongs to.",
        "bom_level": "Depth in the hierarchy (top-level assembly is 0).",
        "unit_of_measure": "Unit for the quantity, e.g. EA, M, KG.",
        "description": "Longer description of the part.",
        "unit_cost": "Cost of one unit.",
        "total_cost": "Quantity multiplied by unit cost.",
        "currency": "ISO 4217 currency code, e.g. USD.",
        "supplier_name": "Name of the supplier.",
        "supplier_part_number": "The supplier's part number.",
        "manufacturer_name": "Name of the manufacturer.",
        "manufacturer_part_number": "The manufacturer's part number (MPN).",
        "material": "Material or material specification.",
        "lead_time_days": "Procurement lead time in days.",
        "reference_designator": "Reference designators, e.g. R1, R2.",
        "notes": "Free-text notes.",
    }
)

_MIN_COLUMN_WIDTH = 12
_COLUMN_PADDING = 4


def template_labels() -> list[str]:
    """The template's header row: every standard field label, required fields first."""
    return [field.label for field in STANDARD_FIELDS]


def build_csv_template() -> bytes:
    """A CSV with one header row, UTF-8 with a byte-order mark (so Excel detects the encoding)
    and CRLF line endings."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(template_labels())
    return buffer.getvalue().encode("utf-8-sig")


def build_xlsx_template() -> bytes:
    """An XLSX workbook: a "BOM" sheet with a bold, frozen header row of the labels, and an
    "Instructions" sheet describing each field."""
    workbook = Workbook()
    sheet = workbook.worksheets[0]  # a new Workbook starts with one worksheet
    sheet.title = BOM_SHEET_TITLE
    bold = Font(bold=True)
    for column, label in enumerate(template_labels(), start=1):
        cell = sheet.cell(row=1, column=column, value=label)
        cell.font = bold
        letter = cell.column_letter
        sheet.column_dimensions[letter].width = max(_MIN_COLUMN_WIDTH, len(label) + _COLUMN_PADDING)
    sheet.freeze_panes = "A2"

    instructions = workbook.create_sheet(INSTRUCTIONS_SHEET_TITLE)
    instructions.append(["Field", "Required", "Description"])
    for cell in instructions[1]:
        cell.font = bold
    for field in STANDARD_FIELDS:
        instructions.append(
            [field.label, "Yes" if field.required else "No", _FIELD_DESCRIPTIONS[field.key]]
        )
    instructions.column_dimensions["A"].width = 28
    instructions.column_dimensions["B"].width = 10
    instructions.column_dimensions["C"].width = 60
    instructions.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
