"""XLSX reading: zip-bomb guards first, then openpyxl in streaming (read-only) mode.

Rows are read through a fixed window of `max_columns + 1` columns, so a row padded out to a
far column is never materialized; a non-empty cell in the extra column fails the parse with
`PARSE_TOO_MANY_COLUMNS`. Cells further right are outside the window and are not seen.
openpyxl fills gaps in the row numbering with empty rows one by one, so the reader stops at
the first row number beyond Excel's 1,048,576, and `build_table` bounds the empty rows.
"""

from __future__ import annotations

import zipfile
import zlib
from collections.abc import Iterator
from pathlib import Path
from typing import IO, TYPE_CHECKING
from xml.etree.ElementTree import Element
from xml.etree.ElementTree import ParseError as XmlParseError

from defusedxml.ElementTree import iterparse as _defused_iterparse  # type: ignore[import-untyped]
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from app.services.parsing.cells import cell_to_text
from app.services.parsing.errors import ParseError, ParseErrorCode
from app.services.parsing.header import (
    EXCEL_MAX_ROWS,
    HeaderScorer,
    RowSource,
    SheetInfo,
    WorkbookTable,
    build_table,
    select_sheet,
    too_complex,
)

if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook

    from app.services.parsing.parse import ParseLimits

OLE_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")
# Small members compress poorly, so their ratio says nothing about a bomb; only members
# larger than this are held to the compression-ratio limit.
_RATIO_CHECK_MIN_BYTES = 1024 * 1024
_ENCRYPTED_FLAG = 0x1

# What openpyxl and zipfile raise for a damaged package. defusedxml's refusals of entity
# tricks are ValueError subclasses.
_CORRUPT_ERRORS = (
    InvalidFileException,
    zipfile.BadZipFile,
    KeyError,
    ValueError,
    TypeError,
    XmlParseError,
    NotImplementedError,
    EOFError,
    zlib.error,
)

# SpreadsheetML namespaces: transitional (what Excel writes) and strict.
_SHEET_NAMESPACES = frozenset(
    {
        "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "http://purl.oclc.org/ooxml/spreadsheetml/main",
    }
)

_CORRUPT_MESSAGE = "The Excel file is damaged or is not a valid .xlsx workbook."
_ENCRYPTED_MESSAGE = "The workbook is password-protected. Remove the password and upload it again."


def read_xlsx(
    path: Path,
    *,
    sheet_name: str | None,
    header_row: int | None,
    limits: ParseLimits,
    header_scorer: HeaderScorer | None = None,
) -> WorkbookTable:
    """Parse one sheet of an .xlsx workbook (cached formula values, never formulas).

    Also counts the chosen sheet's formula cells, from the header row down, that have no cached
    value: they read as empty, and the preview warns about them.
    """
    inspect_archive(path, limits)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except _CORRUPT_ERRORS as exc:
        raise ParseError(ParseErrorCode.FILE_CORRUPT, _CORRUPT_MESSAGE) from exc
    try:
        book, sheet_path = _read_workbook(workbook, sheet_name, header_row, limits, header_scorer)
    finally:
        workbook.close()
    if sheet_path is None:
        return book
    missing = count_formula_cells_without_value(path, sheet_path, book.table.header_row)
    return WorkbookTable(
        sheets=book.sheets,
        sheet_name=book.sheet_name,
        table=book.table,
        formula_cells_without_value=missing,
    )


def count_formula_cells_without_value(path: Path, sheet_path: str, from_row: int) -> int:
    """Cells at or below `from_row` with a formula (`<f>`) but no cached value (`<v>`).

    Streams the sheet's XML member through defusedxml, keeping one row in memory at a time.
    Call only after `inspect_archive`, which bounds the member's expanded size.
    """
    try:
        with zipfile.ZipFile(path) as archive, archive.open(sheet_path.lstrip("/")) as member:
            return _count_in_sheet_xml(member, from_row)
    except _CORRUPT_ERRORS as exc:
        raise ParseError(ParseErrorCode.FILE_CORRUPT, _CORRUPT_MESSAGE) from exc


def _count_in_sheet_xml(source: IO[bytes], from_row: int) -> int:
    count = 0
    row_number = 0
    sheet_data: Element | None = None
    events: Iterator[tuple[str, Element]] = _defused_iterparse(source, events=("start", "end"))
    for event, element in events:
        name = _local_name(element.tag)
        if event == "start":
            if name == "sheetData":
                sheet_data = element
            elif name == "row":
                row_number = _row_number(element.get("r"), row_number + 1)
            continue
        if name == "c" and row_number >= from_row and _lacks_cached_value(element):
            count += 1
        elif name == "row" and sheet_data is not None:
            # Rows are done with once counted; dropping them bounds memory to one row.
            sheet_data.clear()
    return count


def _local_name(tag: str) -> str | None:
    """The element's local name when it is in a SpreadsheetML namespace, else `None`."""
    namespace, _, local = tag[1:].partition("}") if tag.startswith("{") else ("", "", tag)
    return local if namespace in _SHEET_NAMESPACES else None


def _row_number(declared: str | None, implied: int) -> int:
    """The row's `r` attribute, or the row after the previous one when it is absent."""
    if declared is not None and declared.isdigit():
        return int(declared)
    return implied


def _lacks_cached_value(cell: Element) -> bool:
    has_formula = False
    has_value = False
    for child in cell:
        name = _local_name(child.tag)
        if name == "f":
            has_formula = True
        elif name == "v" and child.text:
            has_value = True
    return has_formula and not has_value


def inspect_archive(path: Path, limits: ParseLimits) -> None:
    """Reject archives that are not zips, are encrypted, or would expand unreasonably."""
    with path.open("rb") as handle:
        if handle.read(len(OLE_SIGNATURE)) == OLE_SIGNATURE:
            # Excel stores a password-protected .xlsx inside an OLE compound document.
            raise ParseError(ParseErrorCode.ENCRYPTED_WORKBOOK, _ENCRYPTED_MESSAGE)
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
    except (zipfile.BadZipFile, OSError, EOFError) as exc:
        raise ParseError(ParseErrorCode.FILE_CORRUPT, _CORRUPT_MESSAGE) from exc

    bomb = too_complex("The workbook expands to too much data to process safely.")
    if len(members) > limits.max_zip_members:
        raise bomb
    total = 0
    total_compressed = 0
    for member in members:
        if member.flag_bits & _ENCRYPTED_FLAG:
            raise ParseError(ParseErrorCode.ENCRYPTED_WORKBOOK, _ENCRYPTED_MESSAGE)
        total += member.file_size
        total_compressed += member.compress_size
        if total > limits.max_uncompressed_bytes:
            raise bomb
        if _over_ratio(member.file_size, member.compress_size, limits):
            raise bomb
    # Many members each below the per-member floor can still add up to a bomb.
    if _over_ratio(total, total_compressed, limits):
        raise bomb


def _over_ratio(uncompressed: int, compressed: int, limits: ParseLimits) -> bool:
    return (
        uncompressed > _RATIO_CHECK_MIN_BYTES
        and uncompressed > limits.max_compression_ratio * max(compressed, 1)
    )


def _read_workbook(
    workbook: Workbook,
    sheet_name: str | None,
    header_row: int | None,
    limits: ParseLimits,
    header_scorer: HeaderScorer | None,
) -> tuple[WorkbookTable, str | None]:
    """The chosen sheet's table and the path of its XML member inside the archive."""
    worksheets = {ws.title: ws for ws in workbook.worksheets if isinstance(ws, ReadOnlyWorksheet)}
    sheets = [
        SheetInfo(name=title, visible=ws.sheet_state == "visible")
        for title, ws in worksheets.items()
    ]
    chosen, rows = select_sheet(
        sheets, sheet_name, lambda name: _iter_rows(worksheets[name], limits), limits
    )
    table = build_table(rows, header_row, limits, header_scorer)
    book = WorkbookTable(sheets=sheets[: limits.max_sheets_listed], sheet_name=chosen, table=table)
    # openpyxl keeps the member path of a read-only sheet privately; it is the path its own
    # relationship resolution found for the sheet.
    sheet_path = getattr(worksheets[chosen], "_worksheet_path", None)
    return book, sheet_path if isinstance(sheet_path, str) else None


def _iter_rows(worksheet: ReadOnlyWorksheet, limits: ParseLimits) -> RowSource:
    # Declared dimensions are written by the producing application and may be wrong; reading
    # without them streams every stored row instead of truncating to (or padding up to) them.
    worksheet.reset_dimensions()
    width = limits.max_columns + 1
    # types-openpyxl declares iter_rows for Worksheet only; ReadOnlyWorksheet shares it.
    rows = worksheet.iter_rows(min_col=1, max_col=width, values_only=True)  # type: ignore[misc]
    row_number = 0
    # openpyxl fills gaps with one shared all-empty tuple; recognising it by identity keeps
    # each padding row O(1).
    padding: object = None
    try:
        while True:
            # Sheet XML is parsed lazily, so damage surfaces while iterating.
            try:
                values = next(rows)
            except StopIteration:
                return
            except _CORRUPT_ERRORS as exc:
                raise ParseError(ParseErrorCode.FILE_CORRUPT, _CORRUPT_MESSAGE) from exc
            row_number += 1
            if row_number > EXCEL_MAX_ROWS:
                raise too_complex(f"The sheet has rows beyond row {EXCEL_MAX_ROWS:,}.")
            if values is padding or values.count(None) == len(values):
                padding = values
                yield row_number, ()
                continue
            cells = [cell_to_text(value) for value in values]
            if len(cells) > limits.max_columns and cells[limits.max_columns].strip():
                raise ParseError(
                    ParseErrorCode.TOO_MANY_COLUMNS,
                    f"The file has more than {limits.max_columns:,} columns.",
                )
            yield row_number, cells[: limits.max_columns]
    finally:
        rows.close()
